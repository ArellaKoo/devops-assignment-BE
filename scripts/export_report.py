#!/usr/bin/env python3
"""Assemble the Canvas report, required tables, screenshots and AI appendix.

Run from the workspace (document tooling is separate from app dependencies):
  .local/qwen-tools/.venv/bin/python backend/scripts/export_report.py

Requires python-docx, reportlab and Pillow. The installed tooling environment
already has these packages. No app, database, browser or test is started.
Source Markdown and existing screenshots are read; DOCX/PDF exports are written.
The PDF uses ReportLab directly because LibreOffice is not required.
"""

from __future__ import annotations

import argparse
import html
import re
from dataclasses import dataclass, field
from pathlib import Path

from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Cm, Inches, Pt, RGBColor
from PIL import Image as PILImage
from reportlab.lib import colors
from reportlab.lib.pagesizes import A4
from reportlab.lib.styles import ParagraphStyle, getSampleStyleSheet
from reportlab.lib.units import inch
from reportlab.pdfbase import pdfmetrics
from reportlab.pdfbase.ttfonts import TTFont
from reportlab.platypus import (
    Image, KeepTogether, LongTable, PageBreak, Paragraph, SimpleDocTemplate,
    Spacer, TableStyle,
)


BACKEND = Path(__file__).resolve().parents[1]
WORKSPACE = BACKEND.parent
REPORT_DIR = BACKEND / "docs" / "report"
BLUE = "174469"
REPO_URL = "https://github.com/ArellaKoo/devops-assignment-BE/blob/setup/lab-adaptation/"
FRONTEND_URL = "https://github.com/ArellaKoo/devops-assignment-FE/blob/setup/lab-adaptation/"


@dataclass
class Block:
    kind: str
    text: str = ""
    level: int = 0
    rows: list[list[str]] = field(default_factory=list)
    path: Path | None = None
    source: Path | None = None


def table_cells(line: str) -> list[str]:
    """Split Markdown pipes while preserving pipes inside inline code."""
    placeholders: dict[str, str] = {}

    def protect(match):
        key = f"@@CODE{len(placeholders)}@@"
        placeholders[key] = match.group(0)
        return key

    protected = re.sub(r"`[^`]*`", protect, line.strip().strip("|"))
    cells = re.split(r"(?<!\\)\|", protected)
    for index, cell in enumerate(cells):
        for key, value in placeholders.items():
            cell = cell.replace(key, value)
        cells[index] = cell.strip().replace(r"\|", "|")
    return cells


def parse_markdown(text: str, source: Path | None = None) -> list[Block]:
    lines = text.splitlines()
    blocks: list[Block] = []
    index = 0
    while index < len(lines):
        line = lines[index].strip()
        if not line or re.fullmatch(r"[-*_]{3,}", line):
            index += 1
            continue
        if line.startswith("```"):
            code = []
            index += 1
            while index < len(lines) and not lines[index].strip().startswith("```"):
                code.append(lines[index])
                index += 1
            blocks.append(Block("code", "\n".join(code), source=source))
            index += 1
            continue
        heading = re.match(r"^(#{1,6})\s+(.+)$", line)
        if heading:
            blocks.append(Block("heading", heading[2], len(heading[1]), source=source))
            index += 1
            continue
        if line.startswith("|") and index + 1 < len(lines) and re.match(
            r"^\s*\|?\s*:?-{3,}", lines[index + 1]
        ):
            rows = [table_cells(line)]
            index += 2
            while index < len(lines) and lines[index].strip().startswith("|"):
                rows.append(table_cells(lines[index]))
                index += 1
            width = len(rows[0])
            if any(len(row) != width for row in rows):
                raise ValueError(f"Unequal table columns in {source}: {rows}")
            blocks.append(Block("table", rows=rows, source=source))
            continue
        bullet = re.match(r"^(?:[-*+] |\d+\. )(.+)$", line)
        if bullet:
            fragments = [bullet[1]]
            index += 1
            while index < len(lines):
                nxt = lines[index].strip()
                if not nxt or re.match(r"^(?:#|```|\||[-*+] |\d+\. )", nxt):
                    break
                fragments.append(nxt)
                index += 1
            blocks.append(Block("bullet", " ".join(fragments), source=source))
            continue
        paragraphs = [line.removeprefix("> ")]
        index += 1
        while index < len(lines):
            nxt = lines[index].strip()
            if not nxt or re.match(r"^(?:#|```|\||[-*+] |\d+\. )", nxt):
                break
            paragraphs.append(nxt.removeprefix("> "))
            index += 1
        blocks.append(Block("paragraph", " ".join(paragraphs), source=source))
    return blocks


def markdown_file(path: Path, title: str | None = None) -> list[Block]:
    blocks = parse_markdown(path.read_text(encoding="utf-8"), path)
    if title:
        blocks.insert(0, Block("pagebreak"))
        blocks.insert(1, Block("heading", title, 1))
        if len(blocks) > 2 and blocks[2].kind == "heading":
            blocks.pop(2)
    return blocks


def architecture_appendix() -> list[Block]:
    path = REPORT_DIR / "frontend-architecture.md"
    blocks = markdown_file(path, "Appendix E — Q5(a) routes and shared decisions")
    # Older working notes carry one vertical table per decision. Normalise
    # them to the assessment's one-row-per-decision, four-column layout.
    decisions = []
    kept = []
    current_heading: Block | None = None
    for block in blocks:
        if block.kind == "heading" and re.match(r"\d\. ", block.text):
            current_heading = block
            continue
        if current_heading and block.kind == "table" and len(block.rows[0]) == 2:
            values = {
                re.sub(r"[*`]+", "", row[0]).strip(): row[1]
                for row in block.rows[1:]
            }
            limitation = next(
                (v for k, v in values.items() if k.startswith("What it does not")), ""
            )
            assessment_labels = {
                "1": "Resolving the API's base URL",
                "2": "Attaching the token to a request",
                "3": "Answering an API refusal and where the persona sees it",
                "4": "Gating what a persona can reach and is offered",
                "5": "Request-in-flight duplicate submission control",
            }
            decisions.append([
                assessment_labels.get(current_heading.text[0], current_heading.text.split(" — ", 1)[-1]),
                values.get("Module", ""),
                values.get("Why", "") + " Alternative: " + values.get("Alternative weighed", ""),
                limitation,
            ])
            current_heading = None
            continue
        if current_heading:
            kept.append(current_heading)
            current_heading = None
        kept.append(block)
    if decisions:
        decision_table = Block("table", rows=[[
            "Decision", "Where it lives — module/component",
            "Why there, against the alternative", "What it does not give you",
        ], *decisions], source=path)
        insertion = next(
            (i + 1 for i, b in enumerate(kept) if b.kind == "heading" and b.text == "Five shared decisions"),
            len(kept),
        )
        kept.insert(insertion, decision_table)
    return kept


def screenshot_appendix() -> list[Block]:
    ledger = parse_markdown((REPORT_DIR / "q5-flows.md").read_text(), REPORT_DIR / "q5-flows.md")
    captions = {}
    for block in ledger:
        if block.kind == "table" and len(block.rows[0]) == 2:
            for row in block.rows[1:]:
                if row[0].endswith(".png"):
                    captions[row[0]] = row[1]
    directory = BACKEND / "docs" / "evidence" / "screenshots"
    def order(path):
        match = re.match(r"task(\d+)-(\d+)", path.name)
        return (int(match[1]), int(match[2])) if match else (0, 0)
    images = sorted(directory.glob("*.png"), key=order)
    # Include a separate login capture if one has been added to evidence.
    login_images = sorted((BACKEND / "docs" / "evidence").rglob("*login*.png"))
    for path in reversed(login_images):
        if path not in images:
            images.insert(0, path)
    blocks = [Block("pagebreak"), Block("heading", "Appendix F — Persona flow screenshots", 1)]
    for index, path in enumerate(images, 1):
        if index > 1 and index % 2 == 1:
            blocks.append(Block("pagebreak"))
        caption = captions.get(path.name, path.stem.replace("-", " "))
        blocks.append(Block("image", f"Figure F{index}. {caption}\nSource: {path.relative_to(BACKEND)}", path=path))
    if not images:
        raise FileNotFoundError("No authentic persona screenshots found")
    return blocks


def assemble() -> list[Block]:
    blocks = markdown_file(REPORT_DIR / "report.md")
    first_answer = next(
        i for i, block in enumerate(blocks)
        if block.kind == "heading" and re.match(r"Q1\b", block.text)
    )
    blocks.insert(first_answer, Block("pagebreak"))
    blocks.extend(markdown_file(REPORT_DIR / "q1-audit.md", "Appendix A — Q1(a) four capability audits"))
    blocks.extend(markdown_file(REPORT_DIR / "provenance.md", "Appendix B — Q1(b) module provenance"))
    blocks.extend(markdown_file(REPORT_DIR / "seed-coverage.md", "Appendix C — Q2(c) criterion-by-criterion seed coverage"))
    blocks.extend(markdown_file(REPORT_DIR / "q4-extra-story-tests.md", "Appendix D — Q4(c) designed cases, strategy and priority"))
    blocks.extend(architecture_appendix())
    blocks.extend(screenshot_appendix())
    blocks.extend(markdown_file(BACKEND / "docs" / "assessment" / "ai-prompts.md", "Appendix G — AI prompts and verification disclosure"))
    blocks.extend(markdown_file(BACKEND / "docs" / "assessment" / "delegated-prompts.md", "Appendix G continued — Delegated requests and disclosure gaps"))
    blocks.extend(markdown_file(BACKEND / "docs" / "assessment" / "clean-checkout-rehearsal.md", "Appendix H — Clean-checkout rehearsal record"))
    continuation = BACKEND / "docs" / "assessment" / "codex-handover.md"
    if continuation.exists():
        blocks.extend(markdown_file(continuation, "Appendix H continued — Codex continuation and final verification"))
    return blocks


INLINE = re.compile(r"(\*\*.*?\*\*|`[^`]*`|\[[^\]]+\]\([^)]+\)|(?<!\*)\*[^*]+\*(?!\*))")


def tokens(text):
    for part in INLINE.split(text):
        if not part:
            continue
        if part.startswith("**") and part.endswith("**"):
            yield part[2:-2], "bold", None
        elif part.startswith("`") and part.endswith("`"):
            yield part[1:-1], "code", None
        elif part.startswith("*") and part.endswith("*"):
            yield part[1:-1], "italic", None
        elif part.startswith("[") and "](" in part:
            label, target = part[1:].split("](", 1)
            yield label, "link", target[:-1]
        else:
            yield part, "plain", None


def resolved_link(target: str, source: Path | None) -> str:
    if target.startswith(("https://", "http://", "mailto:")):
        return target
    if target.startswith("#"):
        return ""
    local = (source.parent / target).resolve() if source else BACKEND / target
    for base, url in [(BACKEND, REPO_URL), (WORKSPACE / "frontend", FRONTEND_URL)]:
        if local.is_relative_to(base):
            return url + local.relative_to(base).as_posix()
    return ""


def add_inline(paragraph, text, source=None):
    for value, style, target in tokens(text):
        if style == "link" and (url := resolved_link(target, source)):
            link = OxmlElement("w:hyperlink")
            relationship = paragraph.part.relate_to(
                url, "http://schemas.openxmlformats.org/officeDocument/2006/relationships/hyperlink", is_external=True
            )
            link.set(qn("r:id"), relationship)
            run = OxmlElement("w:r")
            props = OxmlElement("w:rPr")
            color = OxmlElement("w:color")
            color.set(qn("w:val"), BLUE)
            props.append(color)
            run.append(props)
            content = OxmlElement("w:t")
            content.text = value
            run.append(content)
            link.append(run)
            paragraph._p.append(link)
        else:
            run = paragraph.add_run(value)
            run.bold = style == "bold"
            run.italic = style == "italic"
            if style == "code":
                run.font.name = "Courier New"
                run.font.size = Pt(9)


def docx_export(blocks, path):
    document = Document()
    section = document.sections[0]
    section.page_width, section.page_height = Cm(21), Cm(29.7)
    section.top_margin = section.bottom_margin = Inches(.65)
    section.left_margin = section.right_margin = Inches(.65)
    normal = document.styles["Normal"]
    normal.font.name, normal.font.size = "Arial", Pt(10.5)
    normal.paragraph_format.space_after = Pt(6)
    normal.paragraph_format.line_spacing = 1.1
    for name in ["Heading 1", "Heading 2", "Heading 3"]:
        document.styles[name].font.color.rgb = RGBColor.from_string(BLUE)
    header = section.header.paragraphs[0]
    header.text = "ICT381 • SkipQ • Report draft"
    header.style = "Caption"
    footer = section.footer.paragraphs[0]
    footer.alignment = WD_ALIGN_PARAGRAPH.CENTER
    footer.add_run("Page ")
    page_field = OxmlElement("w:fldSimple")
    page_field.set(qn("w:instr"), "PAGE")
    footer._p.append(page_field)
    for block in blocks:
        if block.kind == "pagebreak":
            document.add_page_break()
        elif block.kind == "heading":
            p = document.add_paragraph(style=f"Heading {min(block.level, 3)}")
            add_inline(p, block.text, block.source)
        elif block.kind in {"paragraph", "bullet"}:
            p = document.add_paragraph(style="List Bullet" if block.kind == "bullet" else "Normal")
            add_inline(p, block.text, block.source)
        elif block.kind == "code":
            for line in block.text.splitlines() or [""]:
                p = document.add_paragraph()
                p.paragraph_format.space_after = Pt(0)
                p.paragraph_format.line_spacing = 1
                r = p.add_run(line)
                r.font.name, r.font.size = "Courier New", Pt(8)
        elif block.kind == "table":
            table = document.add_table(rows=1, cols=len(block.rows[0]))
            table.style = "Light Shading Accent 1"
            table.autofit = False
            widths = column_fractions(len(block.rows[0]))
            for cell, fraction in zip(table.rows[0].cells, widths):
                cell.width = Inches(6.97 * fraction)
            repeat = OxmlElement("w:tblHeader")
            repeat.set(qn("w:val"), "true")
            table.rows[0]._tr.get_or_add_trPr().append(repeat)
            for index, row in enumerate(block.rows):
                cells = table.rows[0].cells if index == 0 else table.add_row().cells
                for cell, text, fraction in zip(cells, row, widths):
                    cell.width = Inches(6.97 * fraction)
                    p = cell.paragraphs[0]
                    p.paragraph_format.space_after = Pt(3)
                    p.paragraph_format.keep_together = False
                    p.paragraph_format.keep_with_next = False
                    add_inline(p, text, block.source)
                    for run in p.runs:
                        run.font.size = Pt(8.5 if len(row) >= 4 else 9)
                        if index == 0:
                            run.bold = True
            document.add_paragraph().paragraph_format.space_after = Pt(0)
        elif block.kind == "image":
            p = document.add_paragraph()
            p.alignment = WD_ALIGN_PARAGRAPH.CENTER
            p.paragraph_format.keep_with_next = True
            p.add_run().add_picture(str(block.path), width=Inches(6.65))
            caption = document.add_paragraph(block.text, style="Caption")
            caption.paragraph_format.keep_together = True
            caption.paragraph_format.space_after = Pt(12)
    document.core_properties.title = "ICT381 SkipQ TMA — Report Draft"
    document.core_properties.subject = "Assignment narrative, required tables, authentic screenshots and AI disclosure"
    document.save(path)


def column_fractions(count):
    return {
        2: [.30, .70], 3: [.18, .39, .43],
        4: [.14, .23, .36, .27], 5: [.06, .14, .24, .29, .27],
    }.get(count, [1 / count] * count)


def register_fonts():
    candidates = [
        Path("/Library/Fonts/Arial Unicode.ttf"),
        Path("/System/Library/Fonts/Supplemental/Arial.ttf"),
        Path("/usr/share/fonts/truetype/dejavu/DejaVuSans.ttf"),
    ]
    for candidate in candidates:
        if candidate.exists():
            pdfmetrics.registerFont(TTFont("ReportSans", str(candidate)))
            bold = candidate.with_name("Arial Bold.ttf")
            italic = candidate.with_name("Arial Italic.ttf")
            for suffix, font in [("Bold", bold), ("Italic", italic)]:
                pdfmetrics.registerFont(TTFont("ReportSans" + suffix, str(font if font.exists() else candidate)))
            pdfmetrics.registerFontFamily("ReportSans", normal="ReportSans", bold="ReportSansBold", italic="ReportSansItalic", boldItalic="ReportSansBold")
            return "ReportSans"
    return "Helvetica"


def pdf_markup(text, source=None):
    parts = []
    for value, style, target in tokens(text):
        escaped = html.escape(value)
        if style == "bold":
            escaped = f"<b>{escaped}</b>"
        elif style == "italic":
            escaped = f"<i>{escaped}</i>"
        elif style == "code":
            escaped = f'<font name="Courier" size="8">{escaped}</font>'
        elif style == "link" and (url := resolved_link(target, source)):
            escaped = f'<link href="{html.escape(url, quote=True)}" color="#{BLUE}">{escaped}</link>'
        parts.append(escaped)
    return "".join(parts).replace("\n", "<br/>")


def pdf_export(blocks, path):
    font = register_fonts()
    styles = getSampleStyleSheet()
    styles.add(ParagraphStyle("ReportBody", fontName=font, fontSize=10.2, leading=14, spaceAfter=6))
    styles.add(ParagraphStyle("ReportCell", fontName=font, fontSize=8, leading=10.5, spaceAfter=2, wordWrap="CJK"))
    styles.add(ParagraphStyle("ReportCode", fontName="Courier", fontSize=7.4, leading=10, spaceAfter=1, wordWrap="CJK"))
    styles.add(ParagraphStyle("ReportCaption", fontName=font, fontSize=8.2, leading=10.5, spaceAfter=12))
    for level in [1, 2, 3]:
        styles.add(ParagraphStyle(f"ReportHeading{level}", fontName=font, fontSize=18 - level * 2, leading=22 - level * 2, textColor=colors.HexColor("#" + BLUE), spaceBefore=12, spaceAfter=8, keepWithNext=True))
    width = A4[0] - 1.3 * inch
    flow = []
    for block in blocks:
        if block.kind == "pagebreak":
            flow.append(PageBreak())
        elif block.kind == "heading":
            flow.append(Paragraph(pdf_markup(block.text, block.source), styles[f"ReportHeading{min(block.level, 3)}"]))
        elif block.kind in {"paragraph", "bullet"}:
            flow.append(Paragraph(pdf_markup(block.text, block.source), styles["ReportBody"], bulletText="•" if block.kind == "bullet" else None))
        elif block.kind == "code":
            for line in block.text.splitlines() or [""]:
                # Break very long exact prompts visually, without changing text.
                flow.append(Paragraph(html.escape(line) or "&nbsp;", styles["ReportCode"]))
        elif block.kind == "table":
            rows = [[Paragraph(pdf_markup(cell, block.source), styles["ReportCell"]) for cell in row] for row in block.rows]
            table = LongTable(rows, colWidths=[width * f for f in column_fractions(len(rows[0]))], repeatRows=1, splitByRow=1, splitInRow=0, hAlign="LEFT")
            table.setStyle(TableStyle([
                ("BACKGROUND", (0, 0), (-1, 0), colors.HexColor("#E7EFF5")),
                ("VALIGN", (0, 0), (-1, -1), "TOP"),
                ("GRID", (0, 0), (-1, -1), .35, colors.HexColor("#ACBAC6")),
                ("LEFTPADDING", (0, 0), (-1, -1), 5),
                ("RIGHTPADDING", (0, 0), (-1, -1), 5),
                ("TOPPADDING", (0, 0), (-1, -1), 5),
                ("BOTTOMPADDING", (0, 0), (-1, -1), 5),
            ]))
            flow.extend([table, Spacer(1, 8)])
        elif block.kind == "image":
            with PILImage.open(block.path) as image:
                image_width, image_height = image.size
            display_width = min(width, 6.65 * inch)
            display_height = display_width * image_height / image_width
            flow.append(KeepTogether([
                Image(str(block.path), width=display_width, height=display_height),
                Paragraph(pdf_markup(block.text), styles["ReportCaption"]),
            ]))
    def footer(canvas, doc):
        canvas.saveState()
        canvas.setFont(font, 8)
        canvas.setFillColor(colors.HexColor("#566575"))
        canvas.drawString(.65 * inch, A4[1] - .4 * inch, "ICT381 • SkipQ • Report draft")
        canvas.drawCentredString(A4[0] / 2, .35 * inch, f"Page {doc.page}")
        canvas.restoreState()
    class OutlinedDocument(SimpleDocTemplate):
        def afterFlowable(self, item):
            if not isinstance(item, Paragraph):
                return
            if item.style.name not in {"ReportHeading1", "ReportHeading2"}:
                return
            key = f"heading-{self.seq.nextf('outline')}"
            level = 0 if item.style.name == "ReportHeading1" else 1
            self.canv.bookmarkPage(key)
            self.canv.addOutlineEntry(item.getPlainText(), key, level=level, closed=False)

    document = OutlinedDocument(str(path), pagesize=A4, leftMargin=.65 * inch, rightMargin=.65 * inch, topMargin=.65 * inch, bottomMargin=.65 * inch, title="ICT381 SkipQ TMA — Report Draft", author="")
    document.build(flow, onFirstPage=footer, onLaterPages=footer)


def main():
    parser = argparse.ArgumentParser(description=__doc__)
    parser.add_argument("--output-dir", type=Path, default=REPORT_DIR)
    parser.add_argument("--docx-only", action="store_true")
    args = parser.parse_args()
    blocks = assemble()
    args.output_dir.mkdir(parents=True, exist_ok=True)
    docx_path = args.output_dir / "SkipQ_Report_Draft.docx"
    docx_export(blocks, docx_path)
    print(f"DOCX: {docx_path}")
    if not args.docx_only:
        pdf_path = args.output_dir / "SkipQ_Report_Draft.pdf"
        pdf_export(blocks, pdf_path)
        print(f"PDF: {pdf_path}")
    print(f"Embedded figures: {sum(b.kind == 'image' for b in blocks)}")
    print(f"Tables: {sum(b.kind == 'table' for b in blocks)}")
    print("Draft keeps cover/video placeholders and disclosure gaps visible; review before submission.")


if __name__ == "__main__":
    main()
